from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path

src = Path(r'C:\Users\Billjoy Figueiredo\Downloads\Seva_Innovation_Challenge_2026_Presentation_Template (1).pptx')
out = Path(r'C:\Users\Billjoy Figueiredo\Desktop\seva-innov\output\RoadSync_Seva_Innovation_Challenge_2026.pptx')
pres = Presentation(str(src))

def put(slide_no, shape_i, text):
    shape = pres.slides[slide_no-1].shapes[shape_i]
    tf = shape.text_frame
    # Preserve the template's first paragraph/run formatting while replacing its copy.
    paras = list(tf.paragraphs)
    p = paras[0]
    runs = list(p.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        p.text = text
    for para in paras[1:]:
        para._p.getparent().remove(para._p)
    tf.word_wrap = True

def remove_shape(slide, shape):
    shape._element.getparent().remove(shape._element)

# Cover: remove the template-use guidance box and its contents.
s1 = pres.slides[0]
for sh in list(s1.shapes):
    if sh.left >= Inches(6.35):
        remove_shape(s1, sh)
put(1, 1, 'RoadSync')
put(1, 2, 'Verified road updates for travelers and agencies in Goa')
put(1, 0, 'SEVA INNOVATION CHALLENGE 2026  •  PROJECT PITCH')

# Problem: use official Goa figures and describe the coordination gap as the project hypothesis.
put(2, 3, 'Fragmented road updates leave travelers and agencies without one shared view')
put(2, 6, 'What is the problem?')
put(2, 7, 'Incident and roadwork updates can sit across separate agency channels. Road users may not see a verified, timely picture of changing conditions.')
put(2, 9, 'Who is affected, and where?')
put(2, 10, 'Travelers and road agencies in Goa. RoadSync proposes to start with a pilot in Panaji.')
put(2, 12, 'Evidence · Goa 2023–24')
put(2, 13, 'Goa reported 2,831 motor accident cases, 271 deaths and 1,193 injuries in 2023–24. These figures set the road-safety context for a pilot focused on timely road information.')
put(2, 15, 'Current gaps')
put(2, 16, 'The project identifies fragmented reporting and limited shared visibility. The pilot will map current channels and measure reporting delays.')
put(2, 18, '2,831')
put(2, 19, 'motor accident cases reported in Goa, 2023–24')
put(2, 20, 'Why it matters')
put(2, 21, 'A delayed or unverified update gives travelers and response teams less time to make informed decisions.')

# Solution and workflow.
put(3, 3, 'A shared report, review and public update workflow')
put(3, 6, 'RoadSync connects agency reporting to a public road map. Police and PWD submit location-based reports; an authority reviews them before publication. Travelers see approved active incidents and roadwork.')
put(3, 12, 'Agency report with type, place and time')
put(3, 18, 'Authority reviews and publishes approved reports')
put(3, 24, 'Travelers see verified active road updates')
put(3, 28, 'A shared review queue links agency reports to a public map. This prototype uses sample data; live government and emergency connections remain future work.')

# Feasibility and implementation.
put(4, 5, 'Current stage: Working Prototype · local demo')
put(4, 18, 'Technology')
put(4, 19, 'Built: browser-based interface and role workflows. Pilot needs secure hosting, verified agency access, shared storage, audit logs and onboarding.')
put(4, 21, 'Estimated cost')
put(4, 22, 'Not estimated yet. Scope the secure build, hosting, training and support with the pilot agency before setting a grant amount.')
put(4, 24, 'Timeline')
put(4, 25, '3 months: baseline and Panaji pilot. 6 months: review timing measures and improve the workflow. 12 months: decide whether to expand.')
put(4, 27, 'Proposed partner')
put(4, 28, 'Goa Traffic Police for the first Panaji pilot. Participation and authorization are not yet confirmed.')

# Impact: replace unsupported estimates with pilot scope and measurable outcomes.
put(5, 6, 'Panaji')
put(5, 7, 'Proposed pilot area')
put(5, 8, 'First deployment location')
put(5, 10, '3 months')
put(5, 11, 'Proposed pilot duration')
put(5, 12, 'Time-bound field test')
put(5, 14, '2 measures')
put(5, 15, 'Core workflow timings')
put(5, 16, 'Submission-to-verification; approval-to-publication')
put(5, 18, 'Social impact')
put(5, 19, 'The pilot will measure whether verified updates reach road users sooner and gather traveler and agency feedback.')
put(5, 21, 'Sustainability')
put(5, 22, 'Seek a department budget or grant after pilot scope and costs are validated. A long-term operating owner and funding model remain to be agreed.')
put(5, 24, 'Scale-up plan')
put(5, 25, 'Use Panaji findings to refine the workflow. Expand only with agency agreement, secure infrastructure and a costed implementation plan.')

# Solo builder; populate first, then remove unused team/profile placeholders.
put(6, 7, 'BILL111')
put(6, 8, 'Project lead · sole builder')
put(6, 9, 'Product concept, interface, frontend and workflow implementation')
# Give the solo-builder expertise line room to fit the template column.
pres.slides[5].shapes[9].width = Inches(3.75)
pres.slides[5].shapes[9].height = Inches(0.58)
put(6, 30, 'Support we are seeking')
put(6, 32, 'Technical mentorship and guidance')
put(6, 34, 'Pilot with Goa Traffic Police')
put(6, 36, 'Grant support for a scoped pilot')
put(6, 38, 'Baseline and evaluation support')
put(6, 40, 'Testing, privacy and security review')
put(6, 41, 'Help us turn the local prototype into a secure, measured three-month Panaji pilot with Goa Traffic Police.')
put(6, 42, 'Thank you  |  billjoyxavier19@gmail.com  |  +91 9834086959')
pres.slides[5].shapes[42].height = Inches(0.42)
s6 = pres.slides[5]
for sh in list(s6.shapes):
    l, t = sh.left, sh.top
    in_second = Inches(2.95) <= l <= Inches(5.35) and Inches(1.35) <= t <= Inches(2.55)
    in_third = Inches(0.45) <= l <= Inches(2.95) and Inches(2.65) <= t <= Inches(4.05)
    in_fourth = Inches(2.95) <= l <= Inches(5.35) and Inches(2.65) <= t <= Inches(4.05)
    is_mentor = getattr(sh, 'has_text_frame', False) and ('Mentor / faculty guide' in sh.text or sh.text.strip() == 'Photo')
    if in_second or in_third or in_fourth or is_mentor:
        remove_shape(s6, sh)
out.parent.mkdir(parents=True, exist_ok=True)
pres.save(str(out))
print(out)



